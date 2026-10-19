"""群附件读取与分析输入扩展：ATT-01..08（以及 read_group_attachment 工具）。"""
from __future__ import annotations

import json
import re
import threading
import time
from pathlib import Path

import httpx2
import pytest

from services.im_tools_mcp import attachments as att
from services.im_tools_mcp.attachments import ExtractError, extract_file
from services.im_tools_mcp.im_client import ImClient
from services.im_tools_mcp.tools import (
    TOOL_HANDLERS, ToolInputError, fetch_group_analysis_input, read_group_attachment,
)
from tests.test_im_tools_mcp import SETTINGS, FakeIm, _detail, _group, _msg

MARK = "ATT-MARK-7Q"


# ---------------------------------------------------------------- 夹具文件

def _docx(path: Path, text=f"正文 {MARK}") -> Path:
    from docx import Document
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "表格甲"
    table.rows[0].cells[1].text = "表格乙"
    document.save(str(path))
    return path


def _xlsx(path: Path) -> Path:
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "风险表"
    sheet.append(["风险", MARK])
    sheet.append([None, None])
    sheet.append(["等级", 3])
    book.save(str(path))
    return path


def _pptx(path: Path) -> Path:
    from pptx import Presentation
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = f"标题 {MARK}"
    slide.placeholders[1].text = "要点一"
    presentation.save(str(path))
    return path


def _pdf(path: Path, text: str | None = f"PDF {MARK}") -> Path:
    # 手写最小 PDF：有文字时带文字层，text=None 时是只有空白页的“扫描件”
    content = f"BT /F1 18 Tf 20 100 Td ({text}) Tj ET" if text else ""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for index, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(out)
    return path


# ATT-02：docx / xlsx / pptx / pdf 文字层 / txt / md / csv 都能抽取出标记文本。
@pytest.mark.parametrize("name,maker", [
    ("a.docx", _docx), ("a.xlsx", _xlsx), ("a.pptx", _pptx), ("a.pdf", _pdf),
])
def test_ooxml_and_pdf_extraction(tmp_path, name, maker):
    result = extract_file(maker(tmp_path / name))
    assert result["status"] == "ok", result
    assert MARK in result["text"] and result["chars"] == len(result["text"])
    assert result["format"] == name.split(".")[1] and len(result["sha256"]) == 64


def test_docx_keeps_table_text_and_xlsx_skips_empty_rows(tmp_path):
    assert "表格甲\t表格乙" in extract_file(_docx(tmp_path / "t.docx"))["text"]
    text = extract_file(_xlsx(tmp_path / "t.xlsx"))["text"]
    assert text.splitlines() == ["## 风险表", f"风险\t{MARK}", "等级\t3"]


@pytest.mark.parametrize("name,encoding", [("a.txt", "utf-8"), ("a.md", "utf-8-sig"),
                                           ("a.csv", "gb18030")])
def test_text_formats_detect_encoding(tmp_path, name, encoding):
    path = tmp_path / name
    path.write_bytes(f"列1,列2\n值,{MARK}".encode(encoding))
    result = extract_file(path)
    assert result["status"] == "ok" and MARK in result["text"]


# ATT-03：WPS/旧版 Office 经转换器读取；本机没有 soffice 时用替身转换器验证映射，真实转换须在镜像内补跑。
@pytest.mark.parametrize("name,target,maker", [
    ("a.wps", ".docx", _docx), ("a.doc", ".docx", _docx), ("a.et", ".xlsx", _xlsx),
    ("a.xls", ".xlsx", _xlsx), ("a.dps", ".pptx", _pptx), ("a.ppt", ".pptx", _pptx),
])
def test_legacy_formats_go_through_converter(tmp_path, name, target, maker):
    source = tmp_path / name
    source.write_bytes(b"legacy-binary")
    calls = []

    def converter(staged, target_ext, workdir):
        calls.append((staged.name, target_ext))
        return maker(workdir / ("out" + target))

    result = extract_file(source, converter=converter)
    assert result["status"] == "ok" and MARK in result["text"] and result["format"] == name.split(".")[1]
    assert calls == [("input." + name.split(".")[1], target)]


@pytest.mark.skipif(att.shutil.which("soffice") is None, reason="需要在含 LibreOffice 的镜像内补跑")
def test_real_soffice_converts_docx_roundtrip(tmp_path):  # pragma: no cover - 依赖 soffice
    source = _docx(tmp_path / "x.docx")
    converted = att.convert_with_soffice(source, ".docx", tmp_path)
    assert MARK in extract_file(converted)["text"]


# ATT-04：png 不支持；超过 20MB 过大；损坏的 docx 失败；扫描 PDF 按不支持处理（不阻塞）。
def test_unsupported_too_large_and_corrupt(tmp_path):
    png = tmp_path / "a.png"
    png.write_bytes(b"\x89PNG")
    assert extract_file(png)["status"] == "unsupported"
    big = tmp_path / "big.txt"
    with open(big, "wb") as handle:
        handle.truncate(att.MAX_FILE_BYTES + 1)
    too_large = extract_file(big)
    assert too_large["status"] == "too_large" and too_large["text"] == ""
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    failed = extract_file(bad)
    assert failed["status"] == "failed" and failed["error_code"] == "parse_error"
    scanned = extract_file(_pdf(tmp_path / "scan.pdf", text=None))
    assert scanned["status"] == "unsupported" and scanned["error_code"] == "no_text_layer"
    assert extract_file(tmp_path / "missing.txt")["status"] == "failed"


# ATT-05：正文超过 5 万字截断。
def test_text_is_truncated_at_limit(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("字" * 60_000, encoding="utf-8")
    result = extract_file(path)
    assert result["status"] == "ok" and result["truncated"] is True and result["chars"] == 50_000


# ATT-06：soffice 超时 -> failed；临时目录已删除；并发不超过 2。
def test_converter_timeout_fails_and_cleans_workdir(tmp_path, monkeypatch):
    created = []
    real_mkdtemp = att.tempfile.mkdtemp
    monkeypatch.setattr(att.tempfile, "mkdtemp", lambda **kw: created.append(real_mkdtemp(**kw)) or created[-1])
    source = tmp_path / "a.wps"
    source.write_bytes(b"x")

    def timeout_converter(staged, target_ext, workdir):
        raise ExtractError("converter_timeout")

    result = extract_file(source, converter=timeout_converter)
    assert result["status"] == "failed" and result["error_code"] == "converter_timeout"
    assert created and not Path(created[0]).exists()


def test_soffice_concurrency_is_capped_at_two(tmp_path, monkeypatch):
    monkeypatch.setattr(att.shutil, "which", lambda name: "soffice")
    running, peak, lock = 0, 0, threading.Lock()

    def fake_run(command, **kwargs):
        nonlocal running, peak
        with lock:
            running += 1
            peak = max(peak, running)
        time.sleep(0.15)
        outdir = Path(command[command.index("--outdir") + 1])
        (outdir / (Path(command[-1]).stem + ".docx")).write_bytes(b"x")
        with lock:
            running -= 1
        return att.subprocess.CompletedProcess(command, 0, b"", b"")

    monkeypatch.setattr(att.subprocess, "run", fake_run)
    results = []

    def work(index):
        folder = tmp_path / f"w{index}"
        folder.mkdir()
        src = folder / "in.wps"
        src.write_bytes(b"x")
        results.append(att.convert_with_soffice(src, ".docx", folder))

    threads = [threading.Thread(target=work, args=(i,)) for i in range(5)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert len(results) == 5 and peak <= att.SOFFICE_CONCURRENCY


def test_soffice_missing_and_timeout_are_reported(tmp_path, monkeypatch):
    monkeypatch.setattr(att.shutil, "which", lambda name: None)
    with pytest.raises(ExtractError) as missing:
        att.convert_with_soffice(tmp_path / "a.wps", ".docx", tmp_path)
    assert missing.value.code == "converter_unavailable"
    monkeypatch.setattr(att.shutil, "which", lambda name: "soffice")

    def slow(command, **kwargs):
        raise att.subprocess.TimeoutExpired(command, kwargs["timeout"])

    monkeypatch.setattr(att.subprocess, "run", slow)
    with pytest.raises(ExtractError) as timeout:
        att.convert_with_soffice(tmp_path / "a.wps", ".docx", tmp_path)
    assert timeout.value.code == "converter_timeout"


# ATT-08：摘要与文件实际 sha256 一致；CLI 只显示前 200 字。
def test_sha256_and_cli_preview(tmp_path, capsys):
    import hashlib
    path = tmp_path / "s.txt"
    path.write_text("x" * 500, encoding="utf-8")
    assert extract_file(path)["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert att.main([str(path)]) == 0
    printed = json.loads(capsys.readouterr().out)
    assert printed["status"] == "ok" and len(printed["text"]) == 200 and printed["chars"] == 500
    assert att.main([]) == 2


# ---------------------------------------------------------------- IM 侧


class AttachmentIm(FakeIm):
    """在 FakeIm 上补充 /attachments 接口。"""

    def __init__(self, *args, attachments=None, files=None, forbidden=(), **kwargs):
        super().__init__(*args, **kwargs)
        self.attachments = attachments or {}
        self.files = files or {}
        self.forbidden = set(forbidden)
        self.download_calls = 0

    def handler(self, request):
        path = request.url.path
        match = re.fullmatch(r"/attachments/(\d+)(/download)?", path)
        if match:
            attachment_id = int(match.group(1))
            if attachment_id in self.forbidden or attachment_id not in self.attachments:
                return httpx2.Response(404, json={"detail": "附件不存在或无权访问"})
            if match.group(2):
                self.download_calls += 1
                return httpx2.Response(200, content=self.files[attachment_id])
            return httpx2.Response(200, json=self.attachments[attachment_id])
        return super().handler(request)


def _im(fake):
    return ImClient("http://im.test", "tok", timeout_seconds=5,
                    client=httpx2.AsyncClient(transport=httpx2.MockTransport(fake.handler),
                                              base_url="http://im.test"))


def _meta(name, size=10, mime="application/octet-stream"):
    return {"id": 1, "original_name": name, "file_size": size, "mime_type": mime}


pytestmark = pytest.mark.anyio


@pytest.fixture
def anyio_backend():
    return "asyncio"


# ATT-01：分析输入的每条消息带 sender_id；附件带元数据；原有字段不变。
async def test_analysis_input_has_sender_id_and_attachment_metadata():
    fake = AttachmentIm(
        conversations=[_group(1, "群")],
        details={1: _detail([{"user_id": 1001, "nickname": "阿明"}])},
        messages={1: [
            _msg(1, 1001, "2026-08-01T10:00:00+00:00", "看附件"),
            {**_msg(2, 1002, "2026-08-01T10:05:00+00:00", "", "file", file_name="风险.docx"),
             "attachment_ids": [7, 8]},
        ]},
        attachments={7: _meta("风险.docx", 123, "application/docx"), 8: _meta("表.xlsx", 456)},
    )
    async with _im(fake) as im:
        result = await fetch_group_analysis_input(im, {
            "conversation_id": 1, "start": "2026-08-01T00:00:00+00:00",
            "end": "2026-08-02T00:00:00+00:00"}, settings=SETTINGS)
    first, second = result["messages"]
    assert set(first) == {"message_ref", "timestamp", "text", "sender_id"}
    assert first["sender_id"] == 1001 and first["message_ref"] == "m1" and "阿明" in first["text"]
    assert second["sender_id"] == 1002
    assert second["attachments"] == [
        {"attachment_id": 7, "filename": "风险.docx", "mime_type": "application/docx", "size": 123},
        {"attachment_id": 8, "filename": "表.xlsx", "mime_type": "application/octet-stream", "size": 456}]
    assert result["status"] == "complete" and result["counts"]["included"] == 2


async def test_attachment_metadata_failure_keeps_id_only():
    fake = AttachmentIm(
        conversations=[_group(1, "群")], details={1: _detail([{"user_id": 1001, "nickname": "阿明"}])},
        messages={1: [{**_msg(1, 1001, "2026-08-01T10:00:00+00:00", "", "file", file_name="a.docx"),
                       "attachment_ids": [9]}]},
        forbidden={9}, attachments={9: _meta("x")},
    )
    async with _im(fake) as im:
        result = await fetch_group_analysis_input(im, {
            "conversation_id": 1, "start": "2026-08-01T00:00:00+00:00",
            "end": "2026-08-02T00:00:00+00:00"}, settings=SETTINGS)
    assert result["messages"][0]["attachments"] == [
        {"attachment_id": 9, "filename": "a.docx", "mime_type": "", "size": None}]


def _group_fake(files, attachments, forbidden=()):
    return AttachmentIm(
        conversations=[_group(1, "群")], details={1: _detail([{"user_id": 1001, "nickname": "阿明"}])},
        files=files, attachments=attachments, forbidden=forbidden)


# read_group_attachment 成功路径：下载后抽取文本，返回附件 ID 与文件名。
async def test_read_group_attachment_ok(tmp_path):
    doc = _docx(tmp_path / "r.docx")
    fake = _group_fake({7: doc.read_bytes()}, {7: _meta("r.docx", doc.stat().st_size)})
    async with _im(fake) as im:
        result = await read_group_attachment(im, {"conversation_id": 1, "attachment_id": 7}, settings=SETTINGS)
    assert result["status"] == "ok" and MARK in result["text"]
    assert result["attachment_id"] == 7 and result["filename"] == "r.docx"


# ATT-07：附件不属于该群/无权访问（IM 404/403）-> failed。
async def test_read_group_attachment_forbidden():
    fake = _group_fake({}, {}, forbidden={7})
    async with _im(fake) as im:
        result = await read_group_attachment(im, {"conversation_id": 1, "attachment_id": 7}, settings=SETTINGS)
    assert result["status"] == "failed" and result["error_code"] == "attachment_forbidden"
    assert fake.download_calls == 0


# 元数据声明超过 20MB 时不下载；下载流超过限制也会中止。
async def test_read_group_attachment_too_large_is_not_downloaded():
    fake = _group_fake({7: b"x"}, {7: _meta("big.pdf", att.MAX_FILE_BYTES + 1)})
    async with _im(fake) as im:
        result = await read_group_attachment(im, {"conversation_id": 1, "attachment_id": 7}, settings=SETTINGS)
    assert result["status"] == "too_large" and fake.download_calls == 0


async def test_read_group_attachment_streamed_size_limit(monkeypatch):
    monkeypatch.setattr("services.im_tools_mcp.tools.MAX_FILE_BYTES", 10)
    fake = _group_fake({7: b"x" * 11}, {7: _meta("a.txt", 5)})  # 元数据撒谎，下载流超限
    async with _im(fake) as im:
        result = await read_group_attachment(im, {"conversation_id": 1, "attachment_id": 7}, settings=SETTINGS)
    assert result["status"] == "too_large" and fake.download_calls == 1


async def test_read_group_attachment_validates_arguments():
    fake = _group_fake({}, {})
    async with _im(fake) as im:
        for bad in ({"conversation_id": 0, "attachment_id": 1}, {"conversation_id": 1, "attachment_id": True}):
            with pytest.raises(ToolInputError):
                await read_group_attachment(im, bad, settings=SETTINGS)
    assert "read_group_attachment" in TOOL_HANDLERS
