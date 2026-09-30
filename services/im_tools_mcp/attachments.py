"""群附件文本抽取：docx/xlsx/pptx/pdf(文字层)/txt/md/csv，以及 wps/et/dps/doc/xls/ppt（经 LibreOffice 转换）。

返回 {status: ok|unsupported|too_large|failed, text, truncated, chars, sha256, format, error_code?}。
- unsupported / too_large 不阻塞上层（由调用方固定标注“未读取”）；failed 需要重试。
- 每次抽取使用独立临时目录并在结束后删除；LibreOffice 单次 60 秒超时，全局并发 2，不启用宏。
验证入口：python -m services.im_tools_mcp.attachments <本地文件>（正文只显示前 200 字）。
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
from pathlib import Path
from typing import Any

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 50_000
SOFFICE_TIMEOUT_SECONDS = 60
SOFFICE_CONCURRENCY = 2

_soffice_slots = threading.BoundedSemaphore(SOFFICE_CONCURRENCY)

#: 直接读取的格式
_DIRECT = {".docx": "docx", ".xlsx": "xlsx", ".pptx": "pptx", ".pdf": "pdf",
           ".txt": "text", ".md": "text", ".csv": "text"}
#: 需先转换的格式 -> 转换目标扩展名
_CONVERT = {".wps": ".docx", ".doc": ".docx", ".et": ".xlsx", ".xls": ".xlsx",
            ".dps": ".pptx", ".ppt": ".pptx"}


class ExtractError(Exception):
    """抽取失败；code 为稳定错误码。"""

    # 初始化对象状态并建立所需依赖。
    # English: Initialize with a stable error code.
    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


def _result(status: str, *, fmt: str = "", text: str = "", truncated: bool = False,
            sha256: str = "", error_code: str | None = None) -> dict[str, Any]:
    result: dict[str, Any] = {"status": status, "text": text, "truncated": truncated,
                              "chars": len(text), "sha256": sha256, "format": fmt}
    if error_code:
        result["error_code"] = error_code
    return result


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


# ---- 各格式读取器：返回纯文本，失败抛 ExtractError ----

def _read_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    parts: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            if block.text.strip():
                parts.append(block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    parts.append("\t".join(cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    parts.append("\t".join(cells).rstrip("\t"))
        return "\n".join(parts)
    finally:
        workbook.close()


def _read_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        parts.append(f"## 幻灯片 {index}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        parts.append("\t".join(cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ExtractError("encrypted_file")
    text = "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    if not text:
        # 扫描件没有文字层：不属于解析失败，由调用方标注“未读取”
        raise ExtractError("no_text_layer")
    return text


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    from charset_normalizer import from_bytes

    best = from_bytes(raw).best()
    if best is None:
        raise ExtractError("decode_failed")
    return str(best)


_READERS = {"docx": _read_docx, "xlsx": _read_xlsx, "pptx": _read_pptx, "pdf": _read_pdf,
            "text": _read_text}


# 用 LibreOffice 把 WPS/旧版 Office 转成 OOXML；返回转换后的文件路径。
# English: Convert legacy formats to OOXML with LibreOffice.
def convert_with_soffice(source: Path, target_ext: str, workdir: Path, *,
                         timeout: int = SOFFICE_TIMEOUT_SECONDS) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise ExtractError("converter_unavailable")
    outdir = workdir / "converted"
    outdir.mkdir(exist_ok=True)
    # 独立 profile 目录避免并发实例互相锁定；不启用宏（默认 headless 不执行宏）。
    profile = workdir / "lo-profile"
    command = [soffice, f"-env:UserInstallation={profile.resolve().as_uri()}", "--headless",
               "--norestore", "--convert-to", target_ext.lstrip("."), "--outdir", str(outdir),
               str(source)]
    with _soffice_slots:
        try:
            completed = subprocess.run(command, capture_output=True, timeout=timeout, check=False)
        except subprocess.TimeoutExpired as exc:
            raise ExtractError("converter_timeout") from exc
        except OSError as exc:
            raise ExtractError("converter_unavailable") from exc
    converted = outdir / (source.stem + target_ext)
    if completed.returncode != 0 or not converted.exists():
        raise ExtractError("conversion_failed")
    return converted


# 抽取一个本地文件的文本；converter 可注入（测试替身），默认用 LibreOffice。
# English: Extract text from a local file.
def extract_file(path: str | os.PathLike, *, filename: str | None = None,
                 converter=convert_with_soffice) -> dict[str, Any]:
    source = Path(path)
    name = filename or source.name
    extension = Path(name).suffix.lower()
    try:
        size = source.stat().st_size
    except OSError:
        return _result("failed", error_code="file_missing")
    sha256 = _sha256(source)
    if size > MAX_FILE_BYTES:
        return _result("too_large", fmt=extension.lstrip("."), sha256=sha256, error_code="too_large")
    if extension not in _DIRECT and extension not in _CONVERT:
        return _result("unsupported", fmt=extension.lstrip("."), sha256=sha256,
                       error_code="unsupported_format")
    workdir = Path(tempfile.mkdtemp(prefix="attachment-"))
    try:
        target, fmt = source, _DIRECT.get(extension)
        if extension in _CONVERT:
            target_ext = _CONVERT[extension]
            # 转换器只处理工作目录内的副本，不触碰原文件
            staged = workdir / ("input" + extension)
            shutil.copyfile(source, staged)
            target = converter(staged, target_ext, workdir)
            fmt = _DIRECT[target_ext]
        text = _READERS[fmt](Path(target))
    except ExtractError as exc:
        status = "unsupported" if exc.code in {"no_text_layer", "encrypted_file"} else "failed"
        return _result(status, fmt=extension.lstrip("."), sha256=sha256, error_code=exc.code)
    except Exception:
        # 文件损坏或库内部异常：只暴露稳定错误码，不带异常文本（可能含路径）
        return _result("failed", fmt=extension.lstrip("."), sha256=sha256, error_code="parse_error")
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS]
    return _result("ok", fmt=extension.lstrip("."), text=text, truncated=truncated, sha256=sha256)


# 命令行入口：对本地文件执行同一套抽取，正文只显示前 200 字。
# English: CLI preview of the extraction result.
def main(argv: list[str] | None = None) -> int:
    args = sys.argv[1:] if argv is None else argv
    if len(args) != 1:
        print("用法: python -m services.im_tools_mcp.attachments <本地文件>", file=sys.stderr)
        return 2
    result = extract_file(args[0])
    preview = dict(result)
    preview["text"] = result["text"][:200]
    print(json.dumps(preview, ensure_ascii=False, indent=2))
    return 0 if result["status"] in {"ok", "unsupported", "too_large"} else 1


if __name__ == "__main__":
    raise SystemExit(main())
